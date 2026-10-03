"""Daha önce üretilmiş bulmacaların kaydı. İki işe yarar:

1) 'Anchor' çifti (en üstteki 6 harfli + en soldaki 8 harfli kelime): bu ikili
   bulmacanın geri kalanını büyük ölçüde belirler, aynı ikilinin tekrar
   kullanılmasını engellemek aynı/çok benzer bir bulmacanın bir daha
   üretilmesini önler.
2) Kelime kullanım sayıları: her bulmacadaki TÜM kelimeler ve görsel/bayrakla gösterilen kelimeler de kaydedilir.
   Üretici, eşit durumdaki adaylar arasında daha az kullanılmış kelimeleri
   önce dener - böylece bulmacalar zamanla birbirinden farklılaşır.

Kayıt, kelime bankasıyla AYNI Excel dosyasının ayrı bir SAYFASINDA
("Uretilen_Bulmacalar") tutulur. Excel'e yeni kelime/ipucu eklemeniz bu sayfayı
ETKİLEMEZ - siz elle silmediğiniz sürece kalıcı kalır, üretici her
çalıştığında üstüne eklenir.

ÖNEMLİ: Üretici çalışırken Excel dosyasını Excel'de AÇIK BIRAKMAYIN (Windows
açık dosyaya yazmaya izin vermez).
"""
import os
from collections import Counter
from datetime import datetime

import openpyxl

SHEET_NAME = 'Uretilen_Bulmacalar'
HEADERS = ['Row1_6Harfli', 'Col1_8Harfli', 'Tum_Kelimeler', 'Tarih', 'Gorsel_Kelimeler']


def load_registry(xlsx_path):
    """(kullanılmış anchor çiftleri kümesi, {kelime: kullanım sayısı}, {görsel/bayrak kelimesi: kullanım sayısı}) döner."""
    if not os.path.exists(xlsx_path):
        return set(), Counter(), Counter()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            return set(), Counter(), Counter()
        ws = wb[SHEET_NAME]
        anchors = set()
        usage = Counter()
        image_usage = Counter()
        for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            if not row or not row[0] or not row[1]:
                continue
            anchors.add((str(row[0]).strip().upper(), str(row[1]).strip().upper()))
            if len(row) > 2 and row[2]:
                usage.update(str(row[2]).upper().split())
            if len(row) > 4 and row[4]:
                image_usage.update(str(row[4]).upper().split())
        return anchors, usage, image_usage
    finally:
        wb.close()


def load_used_anchors(xlsx_path):
    return load_registry(xlsx_path)[0]


def save_used_anchor(row1_word, col1_word, xlsx_path, words=None, images=None):
    """Yeni üretilen bir bulmacayı kayda EKLER (mevcut kayıtları ve Excel'deki
    diğer sayfaları etkilemez). Önce geçici bir dosyaya yazılıp sonra yerine
    konur - yazma sırasında bir sorun olursa asıl Excel dosyası bozulmaz."""
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        if SHEET_NAME not in wb.sheetnames:
            ws = wb.create_sheet(SHEET_NAME)
            ws.append(HEADERS)
        else:
            ws = wb[SHEET_NAME]
            for i, h in enumerate(HEADERS, start=1):  # eski kayıtlarda eksik başlıkları tamamla
                if ws.cell(row=1, column=i).value is None:
                    ws.cell(row=1, column=i, value=h)
        ws.append([
            row1_word, col1_word,
            ' '.join(sorted(set(words))) if words else None,
            datetime.now().strftime('%Y-%m-%d %H:%M'),
            ' '.join(sorted(set(images))) if images else None,
        ])
        tmp_path = xlsx_path + '.tmp'
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, xlsx_path)
        except PermissionError as e:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
            raise RuntimeError(
                f"'{os.path.basename(xlsx_path)}' dosyasına yazılamadı. Dosya Excel'de "
                'açık olabilir - kapatıp tekrar deneyin.'
            ) from e
    finally:
        wb.close()
